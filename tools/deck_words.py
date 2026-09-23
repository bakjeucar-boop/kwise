r"""덱 벌의 **네 산출물 글자**를 뜨고, 낱말이 몇 벌에 서는지 세고, 두 판을 맞댄다 (S207 4절).

**S212 에 화면 하나에서 넷으로 넓혔다.** 그 앞까지 이 도구는 화면만 봤고, 판마다
Excel·PPT·Word 는 스크래치가 떠서 **그 스크래치가 판이 끝나면 사라졌다** — S210 이
Excel 을 세 시트만 떠 「비율 0」 을 냈고 S211 이 열두 시트로 다시 세어 **19벌 19줄**을
찾았다. 덤프가 좁으면 **「닫혔다」 로 잘못 읽는다.**

판마다 같은 스크래치를 새로 지었다 — S192(실물 어긋남 열여덟) · S196(잣대
열셋) · S200 · S204(자리 열아홉) · S205(자리 열아홉) · S206(다섯 자리) ·
S207(네 자리). 하는 일은 늘 셋이었다.

    뜬다   ``tools\\render_deck.py`` 의 벌마다 앱을 띄워 **화면 · Excel · PPT ·
           Word** 에 그려진 글자를 모은다 (화면은 ``tools\\screen_audit.py`` 의
           :func:`collect`, 나머지 셋은 실물 바이트를 떠서 연다)
    센다   준 낱말이 **어느 벌 몇 자리**에 서는가
    맞댄다 앞 판 스냅과 줄 단위로 대어 **갈린 줄**과 **수치 조각의 차**를 낸다

**소스 리터럴을 찾지 않는다.** 실물에 그려진 요소만 본다.

    .venv\\Scripts\\python.exe tools\\deck_words.py                       벌 목록과 스냅 자리
    .venv\\Scripts\\python.exe tools\\deck_words.py --snap 앞.json         19벌을 떠 담는다
    .venv\\Scripts\\python.exe tools\\deck_words.py --snap 뒤.json --case small-ind-a1
    .venv\\Scripts\\python.exe tools\\deck_words.py --count 도입 후 --count 목표 역률
    .venv\\Scripts\\python.exe tools\\deck_words.py --read 앞.json --count 도입 후
    .venv\\Scripts\\python.exe tools\\deck_words.py --diff 앞.json 뒤.json

19벌을 다 뜨는 데 **2번 PC 에서 16분 남짓** 걸린다 (S212 982.0초 · S213 983.0초 ·
**네 산출물을 다 굽는 값이다**). 1번 PC 값은 아직 없다 — **「5분 남짓 (S207 324.8초)」
은 화면 하나만 뜨던 때 값이라 걷었다** (S213 3-1). ``--case`` 로 좁히면 벌마다
30~55초다. **담아 둔 스냅이 있으면 ``--read`` 로 읽는다** — 그 판은 앱을 안 띄운다
(S209 2절 · 앞 판이 ``--count`` 로 300.3초를 버린 자리다).
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import time
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:  # 형에만 쓴다 — 도구 시작에 streamlit 을 들이지 않는다
    import pandas as pd
    from matplotlib.figure import Figure
    from matplotlib.text import Text
    from streamlit.testing.v1 import AppTest

    from kwise.report.document import DocumentSections

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "tools"))

import render_deck  # noqa: E402
import screen_audit  # noqa: E402

__all__ = ["BLIND", "MONEY", "Diff", "count_words", "diff", "snap", "snap_dir", "text_of"]

#: 금액·수치 조각. **고치는 판마다 「금액이 한 원도 안 움직였다」 를 이것으로 센다.**
MONEY = re.compile(r"-?[\d,]*\d(?:\.\d+)?\s*(?:원/년|만원/년|억원|만원|원|kW|kWh|kWp|%p|%)")

#: 담는 산출물 넷. **줄의 첫 칸이 이것이다.**
OUTPUTS = ("화면", "Excel", "PPT", "Word")

#: **넓힌 뒤에도 남는 구멍.** 뜰 때도 읽을 때도 이 줄을 찍는다 (S212 2절).
#:
#: 사람이 다음 판에 기억해서 찾지 않게 **도구가 스스로 적는다** — S210 이
#: 「비율 0」 을 낸 까닭이 덤프가 세 시트뿐인 것이었는데 그 사실이 어디에도
#: 안 적혀 있어 다음 판이 그 0 을 값으로 믿었다.
#:
#: **S217 에 그림 안 글자를 담았다** — 그 앞까지는 「화면은 축 이름만 · PPT·Word
#: 그림은 통째」 였고 기온 기준선 「연평균」 이 스냅 0곳이었다.
BLIND = (
    "안 담는 것 — Excel 「15분 시계열」 시트(그 벌 칸의 98.5%지만 글자가 아니다) · "
    "화면 그림의 수·날짜 눈금(vega 가 그릴 때 짓는다)과 툴팁(마우스를 올려야 뜬다) · "
    "png 에만 보이는 잘림·겹침(tools\\capture_screen.py 가 본다)"
)

#: 그림 안 줄의 자리 머리. **줄의 둘째 칸이 이것으로 시작하면 그림 안이다** (S217 2절).
#:
#: 그림 밖 줄은 S216 까지의 스냅과 한 줄도 안 달라야 한다 — 맞댈 때 둘을 갈라
#: 따로 대는 까닭이다(:func:`split`).
FIGURE = "그림:"


def inside(row: list[str]) -> bool:
    return len(row) > 1 and row[1].startswith(FIGURE)


def split(data: dict[str, list[list[str]]], want: bool) -> dict[str, list[list[str]]]:
    """그림 안(``want=True``) 또는 밖 줄만 남긴다."""
    return {key: [row for row in rows if inside(row) == want] for key, rows in data.items()}


# ===================================================================== S217 2절 · 그림 안 글자
#
# **화면** — vega-lite 스펙은 글자 이름(``title``)만 싣고 **자료는 따로 온다**
# (``proto.datasets`` 의 Arrow 바이트). 범례 · 이름 눈금 · 조각·값·기준선 라벨은 그
# 자료의 칸 값이라 스펙만 훑으면 0곳이다. 채널마다 그 칸 값을 푼다.
#
# **PPT · Word** — 그림이 다 :func:`kwise.report.figures.render_png` 한 자리를 지난다.
# 거기를 가로채 **구운 뒤의** matplotlib 글자 객체를 뜬다 — 눈금은 그릴 때 지어지므로
# 굽기 전에는 비어 있다.

#: 채널 → 갈래. 여기 없는 채널(theta · tooltip · detail …)은 그림에 글자를 안 낸다.
_CHANNELS = {
    "color": "범례",
    "fill": "범례",
    "stroke": "범례",
    "x": "눈금",
    "y": "눈금",
    "text": "라벨",
}


def _arrow(payload: bytes) -> pd.DataFrame:
    import pyarrow as pa

    frame: pd.DataFrame = pa.ipc.open_stream(payload).read_all().to_pandas()
    return frame


def _views(
    node: object, data: str | None, encoding: dict[str, Any]
) -> Iterable[tuple[str | None, dict[str, Any]]]:
    """층마다 «자료 이름, 물려받은 채널까지 합친 채널»."""
    if not isinstance(node, dict):
        return
    source = node.get("data")
    name = source.get("name", data) if isinstance(source, dict) else data
    merged = {**encoding, **(node.get("encoding") or {})}
    if "layer" in node:
        for child in node["layer"]:
            yield from _views(child, name, merged)
    elif "mark" in node:
        yield name, merged


def _chart_rows(number: int, spec: str, datasets: Iterable[Any]) -> list[list[str]]:
    parsed = json.loads(spec)
    frames = {item.name: _arrow(item.data.data) for item in datasets}
    names = [text for text in screen_audit._chart_labels(spec) if screen_audit.HANGUL.search(text)]
    where = f"{FIGURE}화면{number:02d} {names[0] if names else ''}".rstrip()
    seen: dict[tuple[str, str], None] = {}
    for data, encoding in _views(parsed, None, {}):
        frame = frames.get(data) if data is not None else None
        for channel, kind in _CHANNELS.items():
            spec_ = encoding.get(channel)
            if not isinstance(spec_, dict) or "field" not in spec_:
                continue
            if kind == "범례" and "legend" in spec_ and spec_["legend"] is None:
                continue
            if kind == "눈금" and ("axis" in spec_ and spec_["axis"] is None):
                continue
            if kind == "눈금" and spec_.get("type") not in {"nominal", "ordinal"}:
                continue  # 수 · 날짜 눈금은 vega 가 짓는다 (:data:`BLIND`)
            domain = (spec_.get("scale") or {}).get("domain")
            if kind == "범례" and isinstance(domain, list):
                values = domain
            elif frame is not None:
                column = str(spec_["field"]).replace("\\", "")
                values = list(frame[column].drop_duplicates()) if column in frame else []
            else:
                values = []
            for value in values:
                if isinstance(value, str) and value.strip():
                    seen[(kind, value)] = None
    return [["화면", where, kind, text] for kind, text in seen]


def _screen_figures(app: AppTest) -> list[list[str]]:
    """화면 차트마다 그림 안 글자. **축 이름·제목은 여기 없다** — 그림 밖 줄(``Chart``)이 쥔다."""
    from streamlit.testing.v1.element_tree import Element

    found: list[Any] = []

    def walk(node: object) -> None:
        proto = getattr(node, "proto", None) if isinstance(node, Element) else None
        if proto is not None and isinstance(getattr(proto, "spec", None), str) and proto.spec:
            found.append(proto)
        children = getattr(node, "children", None)
        for child in children.values() if isinstance(children, dict) else children or []:
            walk(child)

    walk(getattr(app, "main", None))
    walk(getattr(app, "sidebar", None))
    return [
        row
        for n, proto in enumerate(found, 1)
        for row in _chart_rows(n, proto.spec, proto.datasets)
    ]


def _figure_texts(figure: Figure) -> Iterable[tuple[str, str]]:
    """구운 matplotlib 그림 하나의 «갈래, 글자». **보이는 것만.**

    **투명한 글자도 안 보이는 글자다** (S228) — S225 가 역률 100 벌 세로 눈금을
    ``labelcolor="none"`` 으로 감췄는데 ``get_visible()`` 만 봐 벌마다 여섯 줄을 담았다.
    """
    from matplotlib.colors import to_rgba
    from matplotlib.text import Annotation

    def shown(text: Text) -> str:
        if not text.get_visible() or to_rgba(text.get_color(), text.get_alpha())[3] == 0:
            return ""
        return str(text.get_text()).strip()

    for text in figure.texts:
        yield "제목", shown(text)
    for axes in figure.axes:
        for loc in ("left", "center", "right"):
            yield "제목", axes.get_title(loc).strip()
        if axes.axison:
            for axis in (axes.xaxis, axes.yaxis):
                if not axis.get_visible():
                    continue
                yield "축 이름", shown(axis.label)
                locs = axis.get_majorticklocs()
                lo, hi = sorted(axis.get_view_interval())
                for tick, loc in zip(axis.get_major_ticks(len(locs)), locs, strict=True):
                    if lo - 1e-9 <= loc <= hi + 1e-9:
                        for label in (tick.label1, tick.label2):
                            yield "눈금", shown(label)
        legend = axes.get_legend()
        if legend is not None:
            for text in legend.get_texts():
                yield "범례", shown(text)
        for text in axes.texts:
            yield ("주석" if isinstance(text, Annotation) else "라벨"), shown(text)
    for legend in figure.legends:
        for text in legend.get_texts():
            yield "범례", shown(text)


class FigureTap:
    """:func:`kwise.report.figures.render_png` 를 가로채 그림 안 글자를 모은다 (S217 2절).

    ``with`` 로 쓴다. 나가면 원래 함수로 돌려 놓는다. **어느 산출물의 그림인지는
    굽는 순간이 아니라 실물이 정한다** (:meth:`rows`) — 수단 그림은 화면이 다시
    그려질 때 구워지고 Word 는 PPT 가 구운 바이트를 다시 쓴다.
    """

    def __init__(self) -> None:
        #: png 바이트 → «그리는 함수 이름, 갈래와 글자»
        self.pngs: dict[bytes, tuple[str, list[tuple[str, str]]]] = {}
        self._original: Any = None

    def __enter__(self) -> FigureTap:
        from kwise.report import figures

        self._original = figures.render_png

        def tapped(figure: Figure) -> bytes:
            payload: bytes = self._original(figure)
            name = sys._getframe(1).f_code.co_name  # 그리는 함수 이름 (``monthly_peak_png``)
            # **접지 않는다** — 도넛 넷에 「경부하 40%」 가 둘 서면 png 에도 둘이다 (S217 1-4).
            texts = [(k, t) for k, t in _figure_texts(figure) if t]
            self.pngs[payload] = (name, texts)
            return payload

        _swap(figures, "render_png", tapped)
        return self

    def __exit__(self, *_exc: object) -> None:
        from kwise.report import figures

        _swap(figures, "render_png", self._original)

    def rows(self, output: str, pictures: list[tuple[str, bytes]]) -> list[list[str]]:
        """실물에 박힌 그림마다 줄. **가로채지 못한 그림도 한 줄로 적는다** — 구멍이다."""
        rows: list[list[str]] = []
        for where, blob in pictures:
            if blob not in self.pngs:
                rows.append([output, f"{FIGURE}{where}", "못 뜬 그림", f"{len(blob):,}바이트"])
                continue
            name, texts = self.pngs[blob]
            rows += [[output, f"{FIGURE}{where} {name}", kind, text] for kind, text in texts]
        return rows


def _deck_pictures(payload: bytes) -> list[tuple[str, bytes]]:
    """PPT 에 박힌 그림 — 자리는 장 번호."""
    from pptx import Presentation

    found: list[tuple[str, bytes]] = []

    def walk(shapes: Iterable[Any], where: str) -> None:
        for shape in shapes:
            if shape.shape_type == 6:  # 묶음
                walk(shape.shapes, where)
            elif shape.shape_type == 13:  # 그림
                found.append((where, shape.image.blob))

    for number, slide in enumerate(Presentation(io.BytesIO(payload)).slides, 1):
        walk(slide.shapes, f"{number:02d}장")
    return found


def _word_pictures(payload: bytes) -> list[tuple[str, bytes]]:
    """Word 에 박힌 그림 — 문서 차례. 자리는 몇째 그림인지."""
    from docx import Document

    document = Document(io.BytesIO(payload))
    parts = document.part.related_parts
    return [
        (f"{number:02d}번째", parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed].blob)
        for number, shape in enumerate(document.inline_shapes, 1)
    ]


def text_of(row: list[str]) -> str:
    """줄의 글자 — **첫 칸(산출물)을 뺀 나머지 전부** (S212 2절).

    Excel 한 행은 칸이 여럿이고 PPT·Word 표도 그렇다. `row[3]` 한 칸만 보면
    **값 칸이 통째로 세기 밖**이라 「금액이 안 움직였다」 를 못 센다.
    """
    return "\t".join(row[1:])


def snap_dir() -> Path:
    raw = os.environ.get("PROJECT_CACHE")
    return (Path(raw) if raw else PROJECT_ROOT / "cache") / "deck_words"


def _cases(picked: list[str] | None) -> tuple[render_deck.Case, ...]:
    if not picked:
        return render_deck.CASES
    return tuple(render_deck.BY_KEY[key] for key in picked)


def _excel_rows(payload: bytes) -> list[list[str]]:
    """Excel **한 행이 한 줄**. 시계열 시트는 뺀다 (:data:`BLIND`)."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(payload), read_only=True)
    return [
        ["Excel", sheet.title, *(str(v) for v in row if v is not None and str(v).strip())]
        for sheet in book.worksheets
        if "시계열" not in sheet.title
        for row in sheet.iter_rows(values_only=True)
    ]


def _deck_rows(payload: bytes) -> list[list[str]]:
    """PPT **한 문단·표 한 행이 한 줄**. 자리는 장 번호(표는 ``3표`` 꼴)다."""
    from pptx import Presentation

    rows: list[list[str]] = []

    def walk(shapes: Iterable[Any], where: str) -> None:
        for shape in shapes:
            if shape.shape_type == 6:  # 묶음
                walk(shape.shapes, where)
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                rows.extend(
                    ["PPT", where, para.text]
                    for para in shape.text_frame.paragraphs
                    if para.text.strip()
                )
            if getattr(shape, "has_table", False) and shape.has_table:
                rows.extend(
                    ["PPT", f"{where}표", *(cell.text for cell in line.cells)]
                    for line in shape.table.rows
                )

    for number, slide in enumerate(Presentation(io.BytesIO(payload)).slides, 1):
        walk(slide.shapes, str(number))
    return rows


def _word_rows(payload: bytes) -> list[list[str]]:
    """Word **한 문단·표 한 행이 한 줄**. 자리는 문단 스타일(절 제목이 여기서 갈린다)."""
    from docx import Document

    document = Document(io.BytesIO(payload))
    rows: list[list[str]] = [
        ["Word", para.style.name if para.style is not None else "", para.text]
        for para in document.paragraphs
        if para.text.strip()
    ]
    rows += [
        ["Word", f"표{number}", *(cell.text for cell in line.cells)]
        for number, table in enumerate(document.tables, 1)
        for line in table.rows
    ]
    return rows


def _swap(module: object, name: str, value: object) -> object:
    """모듈 속성을 갈아 끼우고 **옛 것을 돌려준다.**

    이름을 글자로 받으므로 mypy 가 모듈 속성 배정을 좁히지 않는다 — 직접
    `compare_view.slides_bytes = grab` 이라 적으면 `[assignment]` 와
    `[attr-defined]` 가 넷 난다(S212 5-6).
    """
    original = getattr(module, name)
    setattr(module, name, value)
    return original


def _artifacts(app: AppTest, key: str) -> list[list[str]]:
    """단추를 눌러 **실물 바이트**를 받아 Excel · PPT · Word 줄을 낸다.

    **Word 단추는 화면에서 감췄으므로**(36세션) PPT 가 받는 재료를 가로채 굽는다 —
    ``tests\\test_base_fee_basis_words.py`` 의 픽스처가 쓰는 길과 같다.
    """
    from kwise.report import slides_bytes
    from kwise.report.document import document_bytes
    from kwise.ui.artifacts import ARTIFACT_KEY
    from kwise.ui.views import compare as compare_view

    captured: dict[str, DocumentSections] = {}

    def grab(sections: DocumentSections) -> tuple[bytes, str]:
        captured["sections"] = sections
        return slides_bytes(sections)

    original = _swap(compare_view, "slides_bytes", grab)
    with FigureTap() as tap:
        try:
            app.button(key="build_ppt").click().run()
            app.button(key="build_excel").click().run()
        finally:
            _swap(compare_view, "slides_bytes", original)
        if app.exception:
            print(f"!! {key} 산출물이 죽었다: {app.exception}")
            return []
        word = document_bytes(captured["sections"])[0]
    store = dict(app.session_state[ARTIFACT_KEY])
    ppt = store["ppt"].payload
    rows = _excel_rows(store["excel"].payload) + _deck_rows(ppt) + _word_rows(word)
    return rows + tap.rows("PPT", _deck_pictures(ppt)) + tap.rows("Word", _word_pictures(word))


def snap(picked: list[str] | None = None) -> dict[str, list[list[str]]]:
    """벌마다 앱을 띄워 **네 산출물** 글자를 모은다. 죽은 벌은 건너뛰고 그 사실을 찍는다."""
    out: dict[str, list[list[str]]] = {}
    for case in _cases(picked):
        started = time.time()
        app = render_deck.build_app(case)
        app.run()
        if app.exception:
            print(f"!! {case.key} 화면이 죽었다: {app.exception}")
            continue
        rows = [["화면", ln.where, ln.kind, ln.slot, ln.text] for ln in screen_audit.collect(app)]
        rows += _screen_figures(app)
        rows += _artifacts(app, case.key)
        out[case.key] = rows
        print(f"    {case.key}\t{_몫({case.key: rows})}\t{time.time() - started:,.1f}초")
    return out


def count_words(data: dict[str, list[list[str]]], words: list[str]) -> dict[str, Counter[str]]:
    """낱말마다 «벌 → 그 벌에서 선 자리 수».

    **그림 안 줄은 글자 칸만 센다** (S217 7절). 자리 칸에 차트 제목이 들어 있어 통째로
    세면 그 그림의 줄마다 제목이 한 번씩 세졌다 — 「왼쪽이 절감」 19 → 178.
    """
    return {
        word: Counter(
            {
                key: hits
                for key, rows in data.items()
                if (hits := sum(word in (row[-1] if inside(row) else text_of(row)) for row in rows))
            }
        )
        for word in words
    }


@dataclass(frozen=True)
class Diff:
    """두 스냅을 맞댄 값."""

    맞댄줄: int
    벌별: Counter[str]
    """벌 → 그 벌에서 갈린 줄 수."""
    갈린짝: Counter[tuple[str, str]]
    """(옛 글자, 새 글자) → 그 짝이 선 벌 수."""
    어긋난벌: tuple[str, ...]
    """한쪽에만 있거나 줄 수가 다른 벌. **맞대지 않았다.**"""
    조각앞: Counter[str]
    조각뒤: Counter[str]

    @property
    def 갈린줄(self) -> int:
        return sum(self.벌별.values())


def diff(before: dict[str, list[list[str]]], after: dict[str, list[list[str]]]) -> Diff:
    """두 스냅을 줄 단위로 맞댄다. **줄 수가 다른 벌은 따로 적는다.**"""
    갈린: Counter[tuple[str, str]] = Counter()
    벌별: Counter[str] = Counter()
    어긋난벌: list[str] = []
    맞댄줄 = 0
    for key in sorted(set(before) | set(after)):
        b, a = before.get(key), after.get(key)
        if b is None or a is None or len(b) != len(a):
            어긋난벌.append(key)
            continue
        맞댄줄 += len(b)
        for rb, ra in zip(b, a, strict=True):
            if rb != ra:
                갈린[(text_of(rb), text_of(ra))] += 1
                벌별[key] += 1

    def 조각(data: dict[str, list[list[str]]]) -> Counter[str]:
        counter: Counter[str] = Counter()
        for rows in data.values():
            for row in rows:
                counter.update(MONEY.findall(text_of(row)))
        return counter

    return Diff(맞댄줄, 벌별, 갈린, tuple(어긋난벌), 조각(before), 조각(after))


def _몫(data: dict[str, list[list[str]]]) -> str:
    """줄 합과 산출물별 몫. **넷 가운데 무엇이 빠졌는지 수로 보인다.**"""
    counts: Counter[str] = Counter()
    for rows in data.values():
        counts.update(row[0] for row in rows)
    몫 = " · ".join(f"{name} {counts[name]:,}" for name in OUTPUTS)
    그림 = sum(inside(row) for rows in data.values() for row in rows)
    return f"{sum(counts.values()):,}줄 ({몫} · 그 가운데 그림 안 {그림:,})"


def _at(path: Path) -> Path:
    """스냅 자리에 댄다. **온 경로면 그대로다** (S210 3절).

    ``--snap`` 과 ``--read`` 는 저마다 이 줄을 들고 있었는데 ``--diff`` 만 없어
    담아 둔 이름(``s210_before.json``)을 주면 ``FileNotFoundError`` 로 죽었다.
    **대는 자리를 여기 하나로 둔다** — 넷째 길이 붙어도 저절로 따라온다.
    """
    return path if path.is_absolute() else snap_dir() / path


def _load(path: Path) -> dict[str, list[list[str]]]:
    data: dict[str, list[list[str]]] = json.loads(_at(path).read_text(encoding="utf-8"))
    return data


def main() -> int:
    parser = argparse.ArgumentParser(description="덱 벌 화면 글자를 뜨고 세고 맞댄다")
    parser.add_argument("--snap", type=Path, help="스냅을 담을 파일 (없는 이름이면 새로 만든다)")
    parser.add_argument("--case", action="append", help="벌 이름. 여러 번 줄 수 있다")
    parser.add_argument("--count", action="append", help="셀 낱말. 여러 번 줄 수 있다")
    parser.add_argument("--diff", nargs=2, type=Path, metavar=("앞", "뒤"), help="두 스냅을 맞댄다")
    parser.add_argument("--read", type=Path, help="뜨지 않고 담아 둔 스냅을 읽는다")
    args = parser.parse_args()

    if args.diff:
        before, after = (_load(path) for path in args.diff)
        # **그림 안과 밖을 따로 댄다** (S217 2절). 그림 안 줄이 없는 옛 스냅과 대도
        # 밖 줄은 한 줄씩 맞대진다 — 섞으면 줄 수가 달라 벌이 통째로 「어긋난 벌」 이다.
        figure = diff(split(before, True), split(after, True))
        앞그림, 뒤그림 = (sum(map(len, split(d, True).values())) for d in (before, after))
        print(f"그림 안 — 앞 {앞그림:,}줄 · 뒤 {뒤그림:,}줄", end="")
        print(f" · 맞댄 줄 {figure.맞댄줄:,} · 갈린 줄 {figure.갈린줄}")
        if figure.어긋난벌:
            print(f"    줄 수가 어긋난 벌 {len(figure.어긋난벌)} — 그림 안은 맞대지 않았다")
        report = diff(split(before, False), split(after, False))
        print(f"그림 밖 — 맞댄 줄 {report.맞댄줄:,} · 갈린 줄 {report.갈린줄}")
        if report.어긋난벌:
            print(f"!! 줄 수가 어긋난 벌 — {', '.join(report.어긋난벌)}")
        print("\n갈린 줄이 있는 벌")
        for key, count in report.벌별.most_common():
            print(f"    {key}\t{count}줄")
        print("\n갈린 글자 짝")
        for (old, new), count in report.갈린짝.most_common():
            print(f"    {count:3d}벌  「{old[:110]}」")
            print(f"          → 「{new[:110]}」")
        앞, 뒤 = sum(report.조각앞.values()), sum(report.조각뒤.values())
        print(f"\n수치 조각 — 앞 {앞:,}개 · 뒤 {뒤:,}개")
        print(f"    준 것 {dict(report.조각앞 - report.조각뒤) or '없다'}")
        print(f"    는 것 {dict(report.조각뒤 - report.조각앞) or '없다'}")
        print(f"\n{BLIND}")
        return 0

    if args.snap is None and args.read is None and not args.count:
        # **인자 없이 돌면 그 판의 값이 다 나온다** — 벌 목록과 담는 자리.
        print(f"스냅 자리 — {snap_dir()}")
        print(f"\n덱 {len(render_deck.CASES)}벌")
        for case in render_deck.CASES:
            pct = case.power_factor_pct
            꼬리 = f"역률 {pct:,.1f}%" if pct is not None else "역률 간주"
            print(f"    {case.key:30s} {case.contract_type} · {case.voltage} · {꼬리}")
        found = sorted(snap_dir().glob("*.json")) if snap_dir().is_dir() else []
        print(f"\n담아 둔 스냅 {len(found)}개")
        for path in found:
            print(f"    {path.name}")
        print(f"\n담는 산출물 — {' · '.join(OUTPUTS)}")
        print(BLIND)
        return 0

    # **읽는 길** (S209 2절). ``--count`` 가 늘 19벌을 다시 떠 그 판 소요를
    # 통째로 버렸다 — 읽는 자리가 ``--diff`` 하나뿐이었다.
    if args.read is not None:
        path = _at(args.read)
        data = _load(path)
        print(f"읽었다 — {path}\n합 {len(data)}벌 · {_몫(data)}")
    else:
        started = time.time()
        print(f"뜬다 — {len(_cases(args.case))}벌 · {' · '.join(OUTPUTS)}")
        data = snap(args.case)
        print(f"합 {len(data)}벌 · {_몫(data)} · {time.time() - started:,.1f}초")
    print(BLIND)

    if args.snap is not None:
        path = _at(args.snap)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(data, ensure_ascii=False, indent=0), encoding="utf-8")
        print(f"담았다 — {path}")

    for word, hits in count_words(data, args.count or []).items():
        print(f"\n「{word}」 — {len(hits)}벌 · 자리 {sum(hits.values())}")
        for key, count in hits.most_common():
            print(f"    {key}\t{count}자리")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
